"""4개 브랜드 배달 샵인샵 월 손익 시뮬레이션 (충북대·사창동, 월세 100만원).

모든 가정값은 이 파일 상단에 모여 있다. 숫자를 바꾸고 다시 실행하면 엑셀이 갱신된다.
사용법: python3 profit_sim.py [OUT.xlsx] [FONT_DIR]
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = sys.argv[1] if len(sys.argv) > 1 else "샵인샵_수익시뮬레이션.xlsx"

# ── 배민 수수료 (본사 마진표와 같은 기준: 배민1 최상위 구간) ─────────────
COMMISSION = 0.078      # 중개이용료
PAYMENT_FEE = 0.033     # 결제수수료
DELIVERY_FEE = 3400     # 점주 부담 배달비 (배민클럽 배달팁 0원 기준)
NET_VAT = 0.04          # 부가세 순부담 (매출세액 1/11 - 매입세액·의제매입 공제) 추정

# ── 3대천왕 김치찜 본사 마진표 (판매가, 원가) + 주문 비중 가정 ─────────
KIMCHI_MENU = [  # (메뉴, 판매가, 원가, 주문비중)
    ("[한그릇] 삼겹듬뿍 김치찜", 12000, 3088, 0.20),
    ("1.5인분 (삼겹/스팸/우삼겹 평균)", 14900, 3873, 0.40),
    ("2~3인 세트 평균", 26967, 7958, 0.30),
    ("3~4인 세트 평균", 35167, 11463, 0.10),
]

# ── 브랜드별 단위 가정 ────────────────────────────────────────────────
# 객단가·원가율: 김치찜은 본사표 가중평균, 나머지는 조사한 메뉴 가격 기준 추정
BRANDS = [
    # 이름, 주문 비중, 객단가, 원가율, 포장재/건, 즉시할인/건, 비고
    dict(name="김치찜 (3대천왕)", share=0.35,
         aov=sum(p * w for _, p, _, w in KIMCHI_MENU),
         cost_rate=sum(c * w for _, _, c, w in KIMCHI_MENU) / sum(p * w for _, p, _, w in KIMCHI_MENU),
         pack=900, coupon=1000, note="본사 마진표 가중평균"),
    dict(name="제육덮밥 (제육대가)", share=0.30, aov=15000, cost_rate=0.33, pack=500, coupon=1000,
         note="1인 13,000원 + 사이드"),
    dict(name="닭발 (30년 불멸의 닭발)", share=0.20, aov=28000, cost_rate=0.33, pack=800, coupon=2000,
         note="세트 22,900~35,900원"),
    dict(name="곱도리탕 (진곱도리탕)", share=0.15, aov=32000, cost_rate=0.38, pack=900, coupon=2000,
         note="소 29,900원~, 곱창 원가 높음"),
]

# ── 월 고정비 ─────────────────────────────────────────────────────────
FIXED = [
    ("월세", 1_000_000),
    ("관리비", 100_000),
    ("전기 (냉장·냉동·후드, 21시간 영업)", 440_000),
    ("가스 (화구, 찜·탕 장시간 조리, 21시간 영업)", 350_000),
    ("수도", 80_000),
    ("통신·POS·주문접수 프로그램", 70_000),
    ("음식물·폐기물 처리", 50_000),
    ("화재·영업배상 보험", 50_000),
    ("세무 기장료", 150_000),
    ("배민 광고 (우리가게클릭 등, 브랜드당 15만)", 600_000),
    ("위생·주방 소모품", 150_000),
    ("로열티·시스템 사용료 (미확인분 대비)", 200_000),
    ("장비 수리·유지", 100_000),
]

# ── 인건비 ────────────────────────────────────────────────────────────
MIN_WAGE = 10_320                 # 2026년 최저시급
HOURLY = round(MIN_WAGE * 1.2 * 1.04, -2)   # 주휴수당(약 20%) + 산재·고용보험 일부
OPEN_HOURS = 21                   # 10:00~다음날 07:00
OWNER_HOURS_MONTH = 10 * 26       # 사장 하루 10시간, 월 26일
OWNER_WAGE = 3_000_000            # 사장 노동의 기회비용 (판단용)


def paid_hours(orders_per_day):
    """매장 상시 1인 공백 + 주문량에 따른 피크 인력."""
    gap = OPEN_HOURS * 30 - OWNER_HOURS_MONTH          # 사장이 없는 시간 = 알바 1명
    peak = 0
    if orders_per_day >= 45:
        peak += 5 * 30                                  # 점심 2h + 저녁 3h 2인 체제
    if orders_per_day >= 70:
        peak += 4 * 30                                  # 야식 피크 3인 체제
    return gap + peak


SCENARIOS = [("비관", 30), ("보수", 45), ("기본", 60), ("낙관", 80)]

# ── 초기 투자 (공시·홍보 기준, 미확인 항목은 가정) ─────────────────────
INITIAL = [
    ("점포 보증금", 10_000_000),
    ("주방 설비 (화구·냉장·후드·작업대, 중고 혼용)", 15_000_000),
    ("3대천왕 김치찜 (가맹금·로열티 0, 초도물품 가정)", 1_000_000),
    ("제육대가 창업비용 (인테리어 포함)", 26_950_000),
    ("30년 불멸의 닭발 (부담금 5,830 + 인테리어 1,650 중 공용분 제외 가정)", 58_300_000),
    ("진곱도리탕 창업비용", 17_990_000),
    ("운영 예비비 (3개월 고정비 일부)", 5_000_000),
]


def unit(b):
    price = b["aov"] - b["coupon"]
    food = b["aov"] * b["cost_rate"]
    fees = price * (COMMISSION + PAYMENT_FEE)
    vat = price * NET_VAT
    contrib = price - food - fees - DELIVERY_FEE - b["pack"] - vat
    return dict(price=price, food=food, fees=fees, vat=vat, contrib=contrib, rate=contrib / price)


def simulate(orders_per_day):
    days = 30
    sales = contrib = food = fees = delivery = pack = vat = 0
    per_brand = []
    for b in BRANDS:
        n = orders_per_day * b["share"] * days
        u = unit(b)
        per_brand.append((b["name"], n, u["price"] * n, u["contrib"] * n))
        sales += u["price"] * n
        contrib += u["contrib"] * n
        food += u["food"] * n
        fees += u["fees"] * n
        delivery += DELIVERY_FEE * n
        pack += b["pack"] * n
        vat += u["vat"] * n
    fixed = sum(v for _, v in FIXED)
    hours = paid_hours(orders_per_day)
    labor = hours * HOURLY
    profit = contrib - fixed - labor
    return dict(orders=orders_per_day, sales=sales, food=food, fees=fees, delivery=delivery, pack=pack,
                vat=vat, contrib=contrib, fixed=fixed, hours=hours, labor=labor, profit=profit,
                after_owner=profit - OWNER_WAGE, per_brand=per_brand)


def breakeven(target):
    """월 영업이익이 target 이상이 되는 최소 일 주문수."""
    for n in range(1, 300):
        if simulate(n)["profit"] >= target:
            return n
    return None


WON = '#,##0"원"'
PCT = '0.0"%"'   # 값은 ×100 해서 저장: 서식을 무시하는 뷰어에서도 21.7로 읽힘


def pct(cell, ratio):
    cell.value = round(ratio * 100, 1)
    cell.number_format = PCT


HEAD = PatternFill("solid", fgColor="2AC1BC")
WB = Font(bold=True, color="FFFFFF")
B = Font(bold=True)


def header(ws, row, cols):
    for c, h in enumerate(cols, 1):
        cell = ws.cell(row, c, h)
        cell.fill, cell.font = HEAD, WB
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def main():
    wb = Workbook()

    # 1) 시나리오 손익
    ws = wb.active
    ws.title = "월손익_시나리오"
    results = [simulate(n) for _, n in SCENARIOS]
    header(ws, 1, ["항목"] + [f"{name} ({n}건/일)" for name, n in SCENARIOS])
    lines = [
        ("월 매출 (할인 후)", "sales"), ("식자재 원가", "food"), ("배민 중개·결제수수료", "fees"),
        ("배달비 (점주 부담)", "delivery"), ("포장재", "pack"), ("부가세 순부담", "vat"),
        ("공헌이익 (변동비 차감 후)", "contrib"), ("고정비 합계", "fixed"),
        ("알바 근무시간 (월)", "hours"), ("인건비 (알바)", "labor"),
        ("영업이익 (사장 몫, 세전)", "profit"), (f"사장 인건비 {OWNER_WAGE // 10000}만원 차감 후", "after_owner"),
    ]
    for i, (label, key) in enumerate(lines, 2):
        ws.cell(i, 1, label).font = B if key in ("contrib", "profit", "after_owner") else Font()
        for j, r in enumerate(results, 2):
            cell = ws.cell(i, j, round(r[key]))
            cell.number_format = '#,##0"시간"' if key == "hours" else WON
            if key in ("profit", "after_owner"):
                cell.font = Font(bold=True, color="C0392B" if r[key] < 0 else "0F8F8A")
    row = len(lines) + 2
    ws.cell(row, 1, "공헌이익률").font = B
    for j, r in enumerate(results, 2):
        pct(ws.cell(row, j), r["contrib"] / r["sales"])
    ws.cell(row + 1, 1, "영업이익률").font = B
    for j, r in enumerate(results, 2):
        pct(ws.cell(row + 1, j), r["profit"] / r["sales"])
    ws.column_dimensions["A"].width = 34
    for col in "BCDE":
        ws.column_dimensions[col].width = 18

    # 2) 브랜드별 단위경제
    ws = wb.create_sheet("브랜드별_건당")
    header(ws, 1, ["브랜드", "주문비중", "객단가", "즉시할인", "실결제", "원가", "원가율", "수수료(11.1%)",
                   "배달비", "포장재", "부가세", "건당 공헌이익", "공헌이익률", "비고"])
    for i, b in enumerate(BRANDS, 2):
        u = unit(b)
        vals = [b["name"], b["share"], round(b["aov"]), b["coupon"], round(u["price"]), round(u["food"]),
                b["cost_rate"], round(u["fees"]), DELIVERY_FEE, b["pack"], round(u["vat"]),
                round(u["contrib"]), u["rate"], b["note"]]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(i, j, v)
            if j in (2, 7, 13):
                pct(cell, v)
            elif isinstance(v, (int, float)):
                cell.number_format = WON
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["N"].width = 28

    # 3) 본사 마진표 검증
    ws = wb.create_sheet("본사마진표_검증")
    header(ws, 1, ["메뉴", "판매가", "원가", "배달비", "중개 7.8%", "결제 3.3%", "본사 계산 마진", "마진율",
                   "포장재", "즉시할인 1,000원", "부가세 4%", "현실 마진", "현실 마진율"])
    table = [("[한그릇] 삼겹듬뿍 김치찜", 12000, 3088), ("[1.5인분] 삼겹듬뿍 김치찜", 14900, 3823),
             ("삼겹듬뿍 2~3인", 26900, 8515), ("삼겹듬뿍 3~4인", 35500, 12363),
             ("1.5인분 스팸폭탄", 14900, 3398), ("스팸폭탄 2~3인", 26500, 6245),
             ("스팸폭탄 3~4인", 33500, 8763), ("1.5인분 우삼겹", 14900, 4398),
             ("우삼겹 2~3인", 27500, 9115), ("우삼겹 3~4인", 36500, 13263)]
    tot_p = tot_m = tot_rp = tot_rm = 0
    for i, (name, p, c) in enumerate(table, 2):
        com, pay = round(p * COMMISSION), round(p * PAYMENT_FEE)
        m = p - c - DELIVERY_FEE - com - pay
        rp = p - 1000
        rm = rp - c - DELIVERY_FEE - rp * (COMMISSION + PAYMENT_FEE) - 900 - rp * NET_VAT
        tot_p, tot_m, tot_rp, tot_rm = tot_p + p, tot_m + m, tot_rp + rp, tot_rm + rm
        for j, v in enumerate([name, p, c, DELIVERY_FEE, com, pay, m, m / p, 900, 1000,
                               round(rp * NET_VAT), round(rm), rm / rp], 1):
            cell = ws.cell(i, j, v)
            if j in (8, 13):
                pct(cell, v)
            elif j > 1:
                cell.number_format = WON
    r = len(table) + 2
    ws.cell(r, 1, "매출 가중 평균").font = B
    pct(ws.cell(r, 8), tot_m / tot_p)
    pct(ws.cell(r, 13), tot_rm / tot_rp)
    ws.cell(r + 1, 1, "본사 표기 '평균 마진율 46%'는 표의 단순평균(43.6%)·가중평균(44.9%) 어느 쪽과도 맞지 않음")
    ws.column_dimensions["A"].width = 30

    # 4) 가정·고정비·투자
    ws = wb.create_sheet("가정_고정비_투자")
    ws.cell(1, 1, "월 고정비").font = B
    for i, (k, v) in enumerate(FIXED, 2):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v).number_format = WON
    r = len(FIXED) + 2
    ws.cell(r, 1, "합계").font = B
    ws.cell(r, 2, sum(v for _, v in FIXED)).number_format = WON
    r += 2
    ws.cell(r, 1, "인건비 가정").font = B
    for k, v, fmt in [("2026 최저시급", MIN_WAGE, WON), ("알바 시간당 실비용 (주휴·보험 포함)", HOURLY, WON),
                      ("영업시간 (하루)", OPEN_HOURS, '0"시간"'), ("사장 근무 (월)", OWNER_HOURS_MONTH, '0"시간"'),
                      ("사장 인건비 환산 (판단용)", OWNER_WAGE, WON)]:
        r += 1
        ws.cell(r, 1, k)
        ws.cell(r, 2, v).number_format = fmt
    r += 1
    ws.cell(r, 1, "※ 5인 미만 사업장은 야간·연장 가산수당 의무 없음 (근로기준법 적용 범위)")
    r += 2
    ws.cell(r, 1, "초기 투자").font = B
    for k, v in INITIAL:
        r += 1
        ws.cell(r, 1, k)
        ws.cell(r, 2, v).number_format = WON
    r += 1
    total_init = sum(v for _, v in INITIAL)
    ws.cell(r, 1, "합계").font = B
    ws.cell(r, 2, total_init).number_format = WON
    r += 2
    ws.cell(r, 1, "손익분기 (일 주문수)").font = B
    for k, target in [("영업이익 0원", 0), ("사장 월 300만원 확보", OWNER_WAGE),
                      ("사장 300만 + 투자금 24개월 회수", OWNER_WAGE + total_init / 24)]:
        r += 1
        ws.cell(r, 1, k)
        ws.cell(r, 2, breakeven(target)).number_format = '0"건/일"'
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 16

    import target_sim
    target_sim.build_target_sheets(wb)
    wb.save(OUT)
    font_dir = sys.argv[2] if len(sys.argv) > 2 else None
    if font_dir:   # NotoSansKR-400/700.ttf가 있는 폴더를 주면 시간대별 그래프 PNG도 만든다
        target_sim.draw_hourly_png("시간대별_처리량.png", font_dir)

    # 콘솔 요약
    for b in BRANDS:
        u = unit(b)
        print(f"{b['name']:<22} 객단가 {b['aov']:>7,.0f} 실결제 {u['price']:>7,.0f} 건당공헌 {u['contrib']:>6,.0f} ({u['rate']:.1%})")
    print("fixed", sum(v for _, v in FIXED), "hourly", HOURLY, "initial", total_init)
    for (name, n), res in zip(SCENARIOS, results):
        print(f"{name} {n}건: 매출 {res['sales']:,.0f} 공헌 {res['contrib']:,.0f} 알바 {res['labor']:,.0f}({res['hours']}h) "
              f"영업이익 {res['profit']:,.0f} 사장인건비후 {res['after_owner']:,.0f}")
    print("BEP 0:", breakeven(0), "/ owner300:", breakeven(OWNER_WAGE),
          "/ owner300+24m payback:", breakeven(OWNER_WAGE + total_init / 24))
    print("saved", OUT)


if __name__ == "__main__":
    main()
