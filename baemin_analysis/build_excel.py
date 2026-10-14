"""배민 캡처 데이터(stores_*.csv)로 상권·경쟁사 분석 엑셀을 만든다."""
import csv
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else "stores_20261012.csv"
OUT = sys.argv[2] if len(sys.argv) > 2 else "배민_상권분석_충북대사창동_20261012.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="2AC1BC")
TOP_FILL = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def num(v, cast=int):
    return cast(v) if v not in ("", None) else None


def load(path):
    rows = []
    with open(path, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            r["주문순위"] = int(r["주문순위"])
            r["별점"] = num(r["별점"], float)
            for k in ("리뷰수", "배달시간(분)", "거리(m)", "최소주문(원)"):
                r[k] = num(r[k])
            r["프랜차이즈"] = "Y" if "프랜차이즈" in r["비고"] else "N"
            rows.append(r)
    return rows


def write_table(ws, header, rows, start_row=1, highlight_top=0):
    for c, h in enumerate(header, 1):
        cell = ws.cell(start_row, c, h)
        cell.fill, cell.font = HEAD_FILL, WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(start_row + i, c, v)
            if i <= highlight_top:
                cell.fill = TOP_FILL
            if isinstance(v, (int, float)) and v >= 1000:
                cell.number_format = "#,##0"
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(header))}{start_row + len(rows)}"


def autosize(ws, widths=None):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[letter].width = min(max(8, longest * 1.6), 45)
    for letter, w in (widths or {}).items():
        ws.column_dimensions[letter].width = w


def dist_band(m):
    if m is None:
        return "미확인"
    if m <= 500:
        return "① 500m 이내"
    if m <= 1500:
        return "② 0.5~1.5km"
    if m <= 3000:
        return "③ 1.5~3km"
    return "④ 3km 초과"


STORE_COLS = ["주문순위", "리뷰순위", "가게명", "업종", "별점", "리뷰수", "배달시간(분)",
              "거리(m)", "거리구간", "최소주문(원)", "배달팁", "배달방식", "프랜차이즈", "광고", "비고"]


def store_row(r):
    return [r["주문순위"], r["리뷰순위"], r["가게명"], r["업종"], r["별점"], r["리뷰수"],
            r["배달시간(분)"], r["거리(m)"], dist_band(r["거리(m)"]), r["최소주문(원)"],
            r["배달팁"], r["배달방식"], r["프랜차이즈"], r["광고"], r["비고"]]


def main():
    stores = load(SRC)
    by_review = sorted(stores, key=lambda r: -(r["리뷰수"] or 0))
    for i, r in enumerate(by_review, 1):
        r["리뷰순위"] = i

    wb = Workbook()

    # 1) 요약
    ws = wb.active
    ws.title = "요약"
    n = len(stores)
    total_reviews = sum(r["리뷰수"] or 0 for r in stores)
    fr = [r for r in stores if r["프랜차이즈"] == "Y"]
    near = [r for r in stores if r["거리(m)"] is not None and r["거리(m)"] <= 500]
    lines = [
        ("배민 상권·경쟁사 분석", None),
        ("기준 위치", "청주 서원구 사창동 · 충북대 정문 인근 (배민 배달주소 기준)"),
        ("수집일", "2026-10-12 18:54~18:55"),
        ("수집 방법", "배민 앱 '음식배달 > 홈(전체 업종)' 을 '주문 많은 순'으로 정렬 후 스크롤 캡처 6장을 판독"),
        ("수집 가게 수", n),
        ("총 리뷰 수 합계", total_reviews),
        ("평균 별점", round(sum(r["별점"] for r in stores) / n, 2)),
        ("프랜차이즈 비중", f"{len(fr)}곳 ({len(fr) / n:.0%})"),
        ("500m 이내 가게", f"{len(near)}곳"),
        (None, None),
        ("※ 유의사항", None),
        ("", "배민 목록에는 '최근 주문 수'가 표시되지 않습니다. '주문순위'는 배민 '주문 많은 순' 정렬 순서이며, 주문량의 상대 지표로 사용하세요."),
        ("", "배민 정렬은 거리·광고(울트라콜/오픈리스트) 등이 섞여 있어 실제 주문수와 완전히 일치하지는 않습니다. '광고' 표시 가게는 순위가 끌어올려졌을 수 있습니다."),
        ("", "캡처 사이 경계에서 일부 가게가 누락됐을 수 있습니다 (22위 가게명 일부 잘림, 54~55위 사이 치킨집 1곳 이름 확인 불가)."),
        ("", "리뷰수는 누적 값이라 오래된 가게일수록 유리합니다. 최근 성과는 '주문순위'와 함께 보세요."),
    ]
    for i, (k, v) in enumerate(lines, 1):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
        if k and v is None:
            ws.cell(i, 1).font = Font(bold=True, size=14 if i == 1 else 11)
        elif k:
            ws.cell(i, 1).font = BOLD
        if isinstance(v, int):
            ws.cell(i, 2).number_format = "#,##0"
        ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 110

    # 2) 주문순 / 3) 리뷰순
    for title, data in (("주문많은순", stores), ("리뷰많은순", by_review)):
        w = wb.create_sheet(title)
        write_table(w, STORE_COLS, [store_row(r) for r in data], highlight_top=10)
        autosize(w, {"C": 38, "O": 40})

    # 4) 업종별
    groups = defaultdict(list)
    for r in stores:
        groups[r["업종"]].append(r)
    cat_rows = []
    for cat, rs in groups.items():
        known = [r for r in rs if r["거리(m)"] is not None]
        best = min(rs, key=lambda r: r["주문순위"])
        cat_rows.append([
            cat, len(rs), sum(r["리뷰수"] for r in rs), round(sum(r["리뷰수"] for r in rs) / len(rs)),
            round(sum(r["별점"] for r in rs) / len(rs), 2),
            min(r["주문순위"] for r in rs),
            round(sum(r["주문순위"] for r in rs) / len(rs), 1),
            round(sum(r["최소주문(원)"] for r in known) / len(known)) if known else None,
            sum(1 for r in known if r["거리(m)"] <= 500),
            sum(1 for r in rs if r["프랜차이즈"] == "Y"),
            best["가게명"],
        ])
    cat_rows.sort(key=lambda x: -x[2])
    w = wb.create_sheet("업종별")
    write_table(w, ["업종", "가게수", "리뷰합계", "평균리뷰", "평균별점", "최고 주문순위",
                    "평균 주문순위", "평균 최소주문(원)", "500m 이내", "프랜차이즈 수", "1위 가게"], cat_rows)
    autosize(w, {"K": 38})

    # 5) 거리구간별
    bands = defaultdict(list)
    for r in stores:
        bands[dist_band(r["거리(m)"])].append(r)
    w = wb.create_sheet("거리구간별")
    write_table(w, ["거리구간", "가게수", "리뷰합계", "평균별점", "평균 배달시간(분)", "가게 목록"], [
        [b, len(rs), sum(r["리뷰수"] for r in rs), round(sum(r["별점"] for r in rs) / len(rs), 2),
         round(sum(r["배달시간(분)"] for r in rs if r["배달시간(분)"]) / max(1, sum(1 for r in rs if r["배달시간(분)"]))),
         ", ".join(r["가게명"] for r in sorted(rs, key=lambda r: r["주문순위"]))]
        for b, rs in sorted(bands.items())
    ])
    autosize(w, {"F": 120})

    wb.save(OUT)
    print("saved", OUT)
    for row in cat_rows:
        print(row)


if __name__ == "__main__":
    main()
